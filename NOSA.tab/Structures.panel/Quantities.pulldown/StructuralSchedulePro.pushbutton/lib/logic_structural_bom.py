# -*- coding: utf-8 -*-
"""
Structural BOM v2.0 — Logic

Collects structural element quantities, differentiates steel vs concrete,
supports per-category density estimation and phase filtering.
"""
import math, io, csv
from collections import defaultdict
from Autodesk.Revit import DB
from nosa_utils.revit_helpers import get_id_value, element_name
from nosa_utils import unit_conversion as _uc10
# ── unit conversions ───────────────────────────────────────────────────────────
FT3_M3  = 0.0283168
FT2_M2  = 0.0929030
FT_M    = _uc10.FT_TO_M
FT2MM   = _uc10.FT_TO_MM
STEEL_DENSITY_KG_M3 = 7850.0

# ── structural categories ──────────────────────────────────────────────────────
CATEGORIES = [
    (u'Structural Columns',      'OST_StructuralColumns'),
    (u'Structural Framing',      'OST_StructuralFraming'),
    (u'Structural Foundations',  'OST_StructuralFoundation'),
    (u'Floors',                  'OST_Floors'),
    (u'Structural Walls',        'OST_Walls'),
]

# Default density used for estimated total weight (concrete incl. rebar)
DEFAULT_DENSITIES = {
    u'Structural Columns':     2500.0,
    u'Structural Framing':     2500.0,
    u'Structural Foundations': 2500.0,
    u'Floors':                 2500.0,
    u'Structural Walls':       2500.0,
}

_STEEL_NAME_KEYS = ['steel', 'acero', 's355', 's275', 's235', 's460',
                    's275jr', 's355jr', 'structural steel', 'chs', 'rhs',
                    'hea', 'heb', 'ipe', ' uc', ' ub', ' pfc']
_STEEL_MAT_CLASSES = ['metal', 'steel', 'acero']

# ── helpers ────────────────────────────────────────────────────────────────────



def _pd(el, bip, default=0.0):
    try:
        p = el.get_Parameter(bip)
        if p and p.HasValue:
            v = p.AsDouble()
            return v if v > 0 else default
    except Exception:
        pass
    return default


def _ps(el, bip, default=''):
    try:
        p = el.get_Parameter(bip)
        if p and p.HasValue:
            return p.AsString() or default
    except Exception:
        pass
    return default


def _pi(el, bip, default=0):
    try:
        p = el.get_Parameter(bip)
        if p and p.HasValue:
            return p.AsInteger()
    except Exception:
        pass
    return default


def _collect(doc, bic):
    return list(DB.FilteredElementCollector(doc)
                .OfCategory(bic)
                .WhereElementIsNotElementType()
                .ToElements())


# ── steel detection ────────────────────────────────────────────────────────────

def _is_steel(el, doc):
    """Return True if the element is made of steel (not concrete)."""
    # 1. StructuralMaterialType property on FamilyInstance
    try:
        if isinstance(el, DB.FamilyInstance):
            mt = el.StructuralMaterialType
            if mt == DB.Structure.StructuralMaterialType.Steel:
                return True
    except Exception:
        pass
    # 2. Structural material parameter → material class / name
    try:
        p = el.get_Parameter(DB.BuiltInParameter.STRUCTURAL_MATERIAL_PARAM)
        if p and p.HasValue:
            mat = doc.GetElement(p.AsElementId())
            if mat:
                cls = (getattr(mat, 'MaterialClass', None) or u'').lower()
                if any(k in cls for k in _STEEL_MAT_CLASSES):
                    return True
                name = (mat.Name or u'').lower()
                if any(k in name for k in _STEEL_NAME_KEYS):
                    return True
    except Exception:
        pass
    return False


# ── level & material helpers ───────────────────────────────────────────────────

def get_levels(doc):
    levels = list(DB.FilteredElementCollector(doc).OfClass(DB.Level).ToElements())
    levels.sort(key=lambda lv: lv.Elevation)
    return [(lv.Name, lv.Id) for lv in levels]


def get_phases(doc):
    phases = list(DB.FilteredElementCollector(doc).OfClass(DB.Phase).ToElements())
    return [(ph.Name, get_id_value(ph.Id)) for ph in phases]


def _level_name(el, doc, level_map):
    for bip in [DB.BuiltInParameter.FAMILY_LEVEL_PARAM,
                DB.BuiltInParameter.LEVEL_PARAM,
                DB.BuiltInParameter.SCHEDULE_LEVEL_PARAM,
                DB.BuiltInParameter.WALL_BASE_CONSTRAINT]:
        try:
            p = el.get_Parameter(bip)
            if p and p.HasValue:
                lvid = get_id_value(p.AsElementId())
                if lvid in level_map:
                    return level_map[lvid]
        except Exception:
            pass
    return u'No level'


def _material_name(el, doc):
    try:
        p = el.get_Parameter(DB.BuiltInParameter.STRUCTURAL_MATERIAL_PARAM)
        if p and p.HasValue:
            mat = doc.GetElement(p.AsElementId())
            if mat:
                return mat.Name or u'—'
    except Exception:
        pass
    try:
        for mat_id in el.GetMaterialIds(False):
            mat = doc.GetElement(mat_id)
            if mat:
                return mat.Name or u'—'
    except Exception:
        pass
    return u'—'


def _family_type_names(el, doc):
    family_name = u'—'
    type_name   = u'—'
    try:
        t = doc.GetElement(el.GetTypeId())
        if isinstance(el, DB.FamilyInstance):
            family_name = el.Symbol.FamilyName or u'—'
            type_name   = (element_name(t) or u'—') if t else u'—'
        else:
            family_name = el.Category.Name if el.Category else u'—'
            type_name   = (element_name(t) or u'—') if t else u'—'
    except Exception:
        pass
    return family_name, type_name


# ── quantity extraction ────────────────────────────────────────────────────────

def _volume_m3(el):
    v = _pd(el, DB.BuiltInParameter.HOST_VOLUME_COMPUTED)
    if v > 0:
        return v * FT3_M3
    try:
        length_ft = _pd(el, DB.BuiltInParameter.CURVE_ELEM_LENGTH)
        if length_ft <= 0:
            length_ft = _pd(el, DB.BuiltInParameter.INSTANCE_LENGTH_PARAM)
        if length_ft > 0:
            # Try section area from geometry bounding box as rough estimate
            try:
                bb = el.get_BoundingBox(None)
                if bb:
                    dx = abs(bb.Max.X - bb.Min.X)
                    dy = abs(bb.Max.Y - bb.Min.Y)
                    area_ft2 = dx * dy
                    if area_ft2 > 0:
                        return length_ft * area_ft2 * FT3_M3
            except Exception:
                pass
    except Exception:
        pass
    return 0.0


def _area_m2(el):
    a = _pd(el, DB.BuiltInParameter.HOST_AREA_COMPUTED)
    if a > 0:
        return a * FT2_M2
    return 0.0


def _length_m(el):
    for bip in [DB.BuiltInParameter.CURVE_ELEM_LENGTH,
                DB.BuiltInParameter.INSTANCE_LENGTH_PARAM]:
        v = _pd(el, bip)
        if v > 0:
            return v * FT_M
    return 0.0


# ── rebar weight index ─────────────────────────────────────────────────────────

def _build_rebar_index(doc):
    idx = defaultdict(float)
    try:
        rebar_class = DB.Structure.Rebar
    except AttributeError:
        try:
            rebar_class = DB.Rebar
        except AttributeError:
            return idx

    for rebar in (DB.FilteredElementCollector(doc)
                  .OfClass(rebar_class)
                  .WhereElementIsNotElementType()
                  .ToElements()):
        try:
            host_id = get_id_value(rebar.GetHostId())
            rtype   = doc.GetElement(rebar.GetTypeId())
            dia_mm  = _pd(rtype, DB.BuiltInParameter.REBAR_BAR_DIAMETER) * FT2MM
            if dia_mm <= 0:
                dia_mm = _pd(rebar, DB.BuiltInParameter.REBAR_BAR_DIAMETER) * FT2MM
            if dia_mm <= 0:
                continue

            area_m2_bar = math.pi * (dia_mm / 2000.0) ** 2

            total_len_ft = _pd(rebar, DB.BuiltInParameter.REBAR_TOTAL_LENGTH_PER_TYPE)
            if total_len_ft <= 0:
                bar_len_ft = _pd(rebar, DB.BuiltInParameter.REBAR_ELEM_LENGTH_PER_TYPE)
                if bar_len_ft <= 0:
                    bar_len_ft = _pd(rebar, DB.BuiltInParameter.CURVE_ELEM_LENGTH)
                n_bars = 1
                try:
                    n_bars = max(rebar.NumberOfBarPositions, 1)
                except Exception:
                    n_bars = max(_pi(rebar, DB.BuiltInParameter.REBAR_NUMBER_OF_SETS), 1)
                total_len_ft = bar_len_ft * n_bars

            idx[host_id] += total_len_ft * FT_M * area_m2_bar * STEEL_DENSITY_KG_M3
        except Exception:
            pass
    return idx


# ── BOM row ────────────────────────────────────────────────────────────────────

class BOMRow(object):
    __slots__ = ['group', 'category', 'family', 'type_name', 'level',
                 'material', 'el_type', 'count', 'volume_m3', 'area_m2',
                 'length_m', 'rebar_kg', 'steel_kg', 'est_kg',
                 'is_subtotal', 'el_ids']

    def __init__(self, **kw):
        for s in self.__slots__:
            setattr(self, s, kw.get(s, None))
        if self.el_ids is None:
            self.el_ids = []

    def add(self, vol, area, length, rebar, steel, est):
        self.count     = (self.count     or 0)   + 1
        self.volume_m3 = (self.volume_m3 or 0.0) + vol
        self.area_m2   = (self.area_m2   or 0.0) + area
        self.length_m  = (self.length_m  or 0.0) + length
        self.rebar_kg  = (self.rebar_kg  or 0.0) + rebar
        self.steel_kg  = (self.steel_kg  or 0.0) + steel
        self.est_kg    = (self.est_kg    or 0.0) + est

    @property
    def vol_str(self):
        v = self.volume_m3 or 0
        return u'{:.3f}'.format(v) if v else u'—'

    @property
    def area_str(self):
        v = self.area_m2 or 0
        return u'{:.2f}'.format(v) if v else u'—'

    @property
    def len_str(self):
        v = self.length_m or 0
        return u'{:.2f}'.format(v) if v else u'—'

    @property
    def rebar_str(self):
        v = self.rebar_kg or 0
        return u'{:.1f}'.format(v) if v > 0.01 else u'—'

    @property
    def steel_str(self):
        v = self.steel_kg or 0
        if v <= 0.01:
            return u'—'
        if v >= 1000:
            return u'{:.2f} t'.format(v / 1000.0)
        return u'{:.0f} kg'.format(v)

    @property
    def est_str(self):
        v = self.est_kg or 0
        if v <= 0.01:
            return u'—'
        if v >= 1000:
            return u'{:.2f} t'.format(v / 1000.0)
        return u'{:.0f} kg'.format(v)

    @property
    def count_str(self):
        return str(self.count or 0)

    @property
    def el_type_str(self):
        return self.el_type or u'—'


# ── collection ─────────────────────────────────────────────────────────────────

def collect_bom(doc, options):
    """
    options = {
        'group_by':           'level' | 'category' | 'material',
        'filter_level_id':    int | None,
        'exclude_phase_id':   int | None,   # phase created in this phase → excluded
        'include_rebar':      bool,
        'densities':          {cat_name: kg_m3, ...},  # user-defined per category
    }
    Returns (list_of_BOMRow, totals_BOMRow).
    """
    group_by         = options.get('group_by', 'level')
    filter_lv_id     = options.get('filter_level_id', None)
    exclude_phase_id = options.get('exclude_phase_id', None)
    include_rebar    = options.get('include_rebar', True)
    densities        = options.get('densities', DEFAULT_DENSITIES)

    level_map = {}
    for lv in DB.FilteredElementCollector(doc).OfClass(DB.Level).ToElements():
        level_map[get_id_value(lv.Id)] = lv.Name

    rebar_idx = _build_rebar_index(doc) if include_rebar else {}

    rows_by_key = {}

    def _key(cat_name, fam, typ, level, material, el_type):
        if group_by == 'level':
            g = level
        elif group_by == 'category':
            g = cat_name
        else:
            g = material
        return (g, cat_name, fam, typ, level, material, el_type)

    for cat_name, bic_name in CATEGORIES:
        bic = getattr(DB.BuiltInCategory, bic_name)
        cat_density = densities.get(cat_name, DEFAULT_DENSITIES.get(cat_name, 2500.0))

        for el in _collect(doc, bic):
            try:
                # Phase filter
                if exclude_phase_id is not None:
                    try:
                        p = el.get_Parameter(DB.BuiltInParameter.PHASE_CREATED)
                        if p and p.HasValue:
                            if get_id_value(p.AsElementId()) == exclude_phase_id:
                                continue
                    except Exception:
                        pass

                # Level filter
                if filter_lv_id is not None:
                    lv_match = False
                    for bip in [DB.BuiltInParameter.FAMILY_LEVEL_PARAM,
                                DB.BuiltInParameter.LEVEL_PARAM,
                                DB.BuiltInParameter.WALL_BASE_CONSTRAINT]:
                        try:
                            p = el.get_Parameter(bip)
                            if p and p.HasValue:
                                if get_id_value(p.AsElementId()) == filter_lv_id:
                                    lv_match = True
                                    break
                        except Exception:
                            pass
                    if not lv_match:
                        continue

                steel    = _is_steel(el, doc)
                el_type  = u'Steel' if steel else u'Concrete'
                lv_name  = _level_name(el, doc, level_map)
                mat_name = _material_name(el, doc)
                fam_name, t_name = _family_type_names(el, doc)

                vol    = _volume_m3(el)
                area   = _area_m2(el)
                length = _length_m(el)
                rebar  = rebar_idx.get(get_id_value(el.Id), 0.0)

                # Steel weight from volume × 7850; only for steel elements
                steel_kg = vol * STEEL_DENSITY_KG_M3 if steel else 0.0
                # Estimated total weight from user density
                est_kg   = vol * cat_density if vol > 0 else 0.0

                k = _key(cat_name, fam_name, t_name, lv_name, mat_name, el_type)
                if k not in rows_by_key:
                    rows_by_key[k] = BOMRow(
                        group       = k[0],
                        category    = cat_name,
                        family      = fam_name,
                        type_name   = t_name,
                        level       = lv_name,
                        material    = mat_name,
                        el_type     = el_type,
                        count       = 0,
                        volume_m3   = 0.0,
                        area_m2     = 0.0,
                        length_m    = 0.0,
                        rebar_kg    = 0.0,
                        steel_kg    = 0.0,
                        est_kg      = 0.0,
                        is_subtotal = False,
                    )
                r = rows_by_key[k]
                r.add(vol, area, length, rebar, steel_kg, est_kg)
                r.el_ids.append(get_id_value(el.Id))

            except Exception:
                pass

    all_rows = sorted(rows_by_key.values(),
                      key=lambda r: (r.group or '', r.category or '',
                                     r.el_type or '', r.family or '',
                                     r.type_name or ''))

    result_rows   = []
    current_group = None
    group_buf     = []

    def _flush(group_rows):
        if not group_rows:
            return
        sub = BOMRow(
            group       = group_rows[0].group,
            category    = u'',
            family      = u'',
            type_name   = u'SUBTOTAL',
            level       = u'',
            material    = u'',
            el_type     = u'',
            count       = sum(r.count     or 0   for r in group_rows),
            volume_m3   = sum(r.volume_m3 or 0.0 for r in group_rows),
            area_m2     = sum(r.area_m2   or 0.0 for r in group_rows),
            length_m    = sum(r.length_m  or 0.0 for r in group_rows),
            rebar_kg    = sum(r.rebar_kg  or 0.0 for r in group_rows),
            steel_kg    = sum(r.steel_kg  or 0.0 for r in group_rows),
            est_kg      = sum(r.est_kg    or 0.0 for r in group_rows),
            is_subtotal = True,
        )
        result_rows.append(sub)

    for row in all_rows:
        g = row.group
        if g != current_group:
            _flush(group_buf)
            group_buf     = []
            current_group = g
        result_rows.append(row)
        group_buf.append(row)
    _flush(group_buf)

    totals = BOMRow(
        group       = u'TOTAL',
        category    = u'',
        family      = u'',
        type_name   = u'GRAND TOTAL',
        level       = u'',
        material    = u'',
        el_type     = u'',
        count       = sum(r.count     or 0   for r in all_rows),
        volume_m3   = sum(r.volume_m3 or 0.0 for r in all_rows),
        area_m2     = sum(r.area_m2   or 0.0 for r in all_rows),
        length_m    = sum(r.length_m  or 0.0 for r in all_rows),
        rebar_kg    = sum(r.rebar_kg  or 0.0 for r in all_rows),
        steel_kg    = sum(r.steel_kg  or 0.0 for r in all_rows),
        est_kg      = sum(r.est_kg    or 0.0 for r in all_rows),
        is_subtotal = True,
    )

    return result_rows, totals


# ── export ────────────────────────────────────────────────────────────────────

def export_xlsx(rows, totals, path, project_name=''):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils  import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = u'Structural BOM'

    accent  = 'FF5F00'
    light   = 'FFF0E8'
    sub_bg  = 'FFD5B0'
    hdr_fnt = Font(bold=True, color='FFFFFFFF', name='Century Gothic', size=10)
    sub_fnt = Font(bold=True, name='Century Gothic', size=10)
    tot_fnt = Font(bold=True, color='FFFFFFFF', name='Century Gothic', size=10)
    reg_fnt = Font(name='Century Gothic', size=9)
    center  = Alignment(horizontal='center', vertical='center')
    thin    = Side(border_style='thin', color='FFE0E0E0')
    bdr     = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells('A1:N1')
    ws['A1'] = u'STRUCTURAL BOM — {}'.format(project_name or 'PROJECT')
    ws['A1'].font      = Font(bold=True, size=14, name='Century Gothic', color=accent)
    ws['A1'].alignment = center

    headers = [u'Group', u'Category', u'Family', u'Type', u'Level', u'Material',
               u'Mat. Type', u'No.', u'Vol. (m³)', u'Area (m²)',
               u'Length (m)', u'Rebar (kg)', u'Steel (kg)', u'Est. Wt.']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font      = hdr_fnt
        cell.fill      = PatternFill('solid', fgColor=accent)
        cell.alignment = center
        cell.border    = bdr

    row_num = 4
    for r in rows:
        if r.is_subtotal and r.type_name in (u'SUBTOTAL', u'GRAND TOTAL'):
            values = [r.group, u'', u'', r.type_name, u'', u'', u'',
                      r.count or 0,
                      round(r.volume_m3 or 0, 3),
                      round(r.area_m2   or 0, 2),
                      round(r.length_m  or 0, 2),
                      round(r.rebar_kg  or 0, 1),
                      round(r.steel_kg  or 0, 1),
                      round(r.est_kg    or 0, 1)]
            fill = PatternFill('solid', fgColor=sub_bg)
            font = sub_fnt
        else:
            values = [r.group, r.category, r.family, r.type_name,
                      r.level, r.material, r.el_type or '',
                      r.count or 0,
                      round(r.volume_m3 or 0, 3) or u'',
                      round(r.area_m2   or 0, 2) or u'',
                      round(r.length_m  or 0, 2) or u'',
                      round(r.rebar_kg  or 0, 1) or u'',
                      round(r.steel_kg  or 0, 1) or u'',
                      round(r.est_kg    or 0, 1) or u'']
            fill = PatternFill('solid', fgColor=light) if row_num % 2 == 0 else None
            font = reg_fnt

        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col, value=v)
            cell.font   = font
            cell.border = bdr
            if fill:
                cell.fill = fill
            if col >= 8:
                cell.alignment = center
        row_num += 1

    if totals:
        tot_vals = [u'', u'', u'', u'GRAND TOTAL', u'', u'', u'',
                    totals.count or 0,
                    round(totals.volume_m3 or 0, 3),
                    round(totals.area_m2   or 0, 2),
                    round(totals.length_m  or 0, 2),
                    round(totals.rebar_kg  or 0, 1),
                    round(totals.steel_kg  or 0, 1),
                    round(totals.est_kg    or 0, 1)]
        for col, v in enumerate(tot_vals, 1):
            cell = ws.cell(row=row_num, column=col, value=v)
            cell.font      = tot_fnt
            cell.fill      = PatternFill('solid', fgColor=accent)
            cell.border    = bdr
            cell.alignment = center

    widths = [16, 18, 22, 28, 14, 20, 10, 6, 12, 12, 12, 12, 12, 14]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = 'A4'
    wb.save(path)


def export_csv(rows, totals, path):
    headers = ['Group', 'Category', 'Family', 'Type', 'Level', 'Material',
               'Mat. Type', 'No.', 'Vol.(m³)', 'Area(m²)', 'Length(m)',
               'Rebar(kg)', 'Steel(kg)', 'Est.Wt.']
    with io.open(path, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            w.writerow([
                r.group or '', r.category or '', r.family or '',
                r.type_name or '', r.level or '', r.material or '',
                r.el_type or '',
                r.count or 0,
                round(r.volume_m3 or 0, 3),
                round(r.area_m2   or 0, 2),
                round(r.length_m  or 0, 2),
                round(r.rebar_kg  or 0, 1),
                round(r.steel_kg  or 0, 1),
                round(r.est_kg    or 0, 1),
            ])
        if totals:
            w.writerow(['', '', '', 'GRAND TOTAL', '', '', '',
                        totals.count or 0,
                        round(totals.volume_m3 or 0, 3),
                        round(totals.area_m2   or 0, 2),
                        round(totals.length_m  or 0, 2),
                        round(totals.rebar_kg  or 0, 1),
                        round(totals.steel_kg  or 0, 1),
                        round(totals.est_kg    or 0, 1)])
