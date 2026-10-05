# -*- coding: utf-8 -*-
"""T7.5 BS 8666 shape sketches: fitting and letter placement (no Revit)."""
from __future__ import absolute_import, division, print_function, unicode_literals
import os
import sys

_lib = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'lib'))
if _lib not in sys.path:
    sys.path.insert(0, _lib)

import shape_sketch as sk  # noqa: E402

# Revit's browser sketch of shape 21 (U-bar), mm
SHAPE_21 = [[(8, 0), (8, -292)], [(8, -292), (992, -292)], [(992, -292), (992, 0)]]


def test_fit_keeps_the_sketch_inside_the_margins_and_centred():
    lines, _to_px = sk.fit(SHAPE_21)
    xs = [p[0] for l in lines for p in l]
    ys = [p[1] for l in lines for p in l]
    assert min(xs) >= sk.MARGIN - 1e-6 and max(xs) <= sk.WIDTH - sk.MARGIN + 1e-6
    assert min(ys) >= sk.MARGIN - 1e-6 and max(ys) <= sk.HEIGHT - sk.MARGIN + 1e-6
    assert abs((min(xs) + max(xs)) / 2.0 - sk.WIDTH / 2.0) < 1e-6


def test_the_u_opens_upwards_in_the_image():
    lines, _to_px = sk.fit(SHAPE_21)
    leg_top, leg_bottom = lines[0][0], lines[0][1]
    assert leg_top[1] < leg_bottom[1]                 # image y grows downwards


def test_letters_sit_outside_the_shape():
    lines, _to_px = sk.fit(SHAPE_21)
    labels = sk.label_positions(lines, [u'A', u'B', u'C'])
    (ta, xa, ya), (tb, xb, yb), (tc, xc, yc) = labels
    assert (ta, tb, tc) == (u'A', u'B', u'C')
    assert xa < lines[0][0][0] and xc > lines[2][0][0]   # legs: letters outside, left and right
    assert yb > lines[1][0][1]                           # base: letter below it


def test_straight_bar_fills_the_width():
    lines, _to_px = sk.fit([[(0, 0), (1000, 0)]])
    assert abs(lines[0][0][0] - sk.MARGIN) < 1e-6 and abs(lines[0][1][0] - (sk.WIDTH - sk.MARGIN)) < 1e-6


def test_arc_tessellation_ends_on_its_points():
    pts = sk.tessellate_arc((1.0, 0.0), (0.0, 1.0), (0.0, 0.0), steps=8)
    assert abs(pts[0][0] - 1.0) < 1e-9 and abs(pts[-1][1] - 1.0) < 1e-9 and len(pts) == 9


def _decode(png):
    import struct
    import zlib
    assert png[:8] == b'\x89PNG\r\n\x1a\n'
    pos, chunks = 8, {}
    while pos < len(png):
        length = struct.unpack('>I', png[pos:pos + 4])[0]
        kind = png[pos + 4:pos + 8]
        chunks[kind] = png[pos + 8:pos + 8 + length]
        pos += 12 + length
    width, height = struct.unpack('>II', chunks[b'IHDR'][:8])
    raw = zlib.decompress(chunks[b'IDAT'])
    rows = [raw[r * (width + 1) + 1:(r + 1) * (width + 1)] for r in range(height)]
    return width, height, rows, chunks


def test_png_sketch_draws_the_u_and_its_letters():
    import png_sketch
    lines, _to_px = sk.fit(SHAPE_21)
    labels = sk.label_positions(lines, [u'A', u'B', u'C'])
    png = png_sketch.draw_sketch(lines, labels, sk.WIDTH, sk.HEIGHT).png(dpi=sk.DPI)
    width, height, rows, chunks = _decode(png)
    assert (width, height) == (sk.WIDTH, sk.HEIGHT)
    assert chunks[b'pHYs'][:4] == (35433).to_bytes(4, 'big')          # 900 dpi: 30 x 15 mm, 3x sharper
    x, y = [int(round(v)) for v in lines[1][0]]                          # a corner of the U: inked
    assert rows[y][x] < 60
    assert rows[2][2] == 255                                             # paper elsewhere
    _t, lx, ly = labels[1]
    half = 13 * sk.SCALE
    letter = [rows[yy][xx] for yy in range(int(ly) - half, int(ly) + half)
              for xx in range(int(lx) - half, int(lx) + half)]
    assert min(letter) == 0                                              # the 'B' is drawn


def test_bbs_excel_carries_the_sketch_beside_the_shape_code():
    import tempfile
    import zipfile
    _ext = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..', '..', 'lib'))
    if _ext not in sys.path:
        sys.path.insert(0, _ext)
    import png_sketch
    import rebar_schedule
    folder = tempfile.mkdtemp()
    lines, _to_px = sk.fit(SHAPE_21)
    png_path = os.path.join(folder, 'NOSA_BS8666_21.png')
    with open(png_path, 'wb') as f:
        f.write(png_sketch.draw_sketch(lines, sk.label_positions(lines, ['A', 'B', 'C']),
                                       sk.WIDTH, sk.HEIGHT).png())
    data = [{'member': u'C1', 'mark': u'01', 'diameter_mm': 12, 'count': 6, 'members': 1,
             'unit_length_mm': 1500.0, 'shape_code': u'21', 'shape_params': u'A=300;B=900;C=300',
             'total_length_mm': 9000.0, 'total_weight_kg': 8.0, 'revision': u'A', 'shape_id': 77}]
    out = os.path.join(folder, 'bbs.xlsx')
    assert rebar_schedule.export_xlsx(data, out, {77: png_path})
    z = zipfile.ZipFile(out)
    assert 'xl/media/image1.png' in z.namelist()
    z.close()
    try:
        import openpyxl
    except ImportError:
        return
    ws = openpyxl.load_workbook(out).active
    assert ws['H1'].value == u'Shape code' and ws['I1'].value == u'Shape'
    assert ws['H2'].value == u'21' and ws['B2'].value == u'01' and ws['J2'].value == 300
    assert ws['R1'].value == u'Rev.' and ws['R2'].value == u'A' and ws['Q2'].value == 8.0
