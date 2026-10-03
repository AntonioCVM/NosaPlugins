# -*- coding: utf-8 -*-
"""nosa_utils.xlsx_writer: a valid workbook with text, numbers and a picture (T7.5)."""
import os
import sys
import tempfile
import unittest
import zipfile

_LIB = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'lib'))
if _LIB not in sys.path:
    sys.path.insert(0, _LIB)

from nosa_utils.xlsx_writer import column_letter, write_xlsx  # noqa: E402

# 1x1 white PNG
_PNG = (b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x00\x00\x00\x00:~\x9bU'
        b'\x00\x00\x00\nIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82')


class XlsxWriterTests(unittest.TestCase):

    def test_column_letters(self):
        self.assertEqual([column_letter(i) for i in (0, 25, 26, 27, 701, 702)],
                         ['A', 'Z', 'AA', 'AB', 'ZZ', 'AAA'])

    def test_workbook_parts_and_values(self):
        path = os.path.join(tempfile.mkdtemp(), 'bbs.xlsx')
        write_xlsx(path, u'Bar Bending Schedule', [[u'Member', u'Bar mark', u'Weight [kg]'],
                                                   [u'C1 & C2', u'01', 12.5]],
                   widths=[12, 8, 10], heights={1: 45}, images=[{'row': 1, 'col': 1, 'png': _PNG,
                                                                 'size_mm': (30, 15)}])
        z = zipfile.ZipFile(path)
        names = set(z.namelist())
        z.close()
        for part in ('[Content_Types].xml', 'xl/workbook.xml', 'xl/worksheets/sheet1.xml',
                     'xl/drawings/drawing1.xml', 'xl/media/image1.png'):
            self.assertIn(part, names)
        try:
            import openpyxl
        except ImportError:
            return
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        self.assertEqual(ws.title, u'Bar Bending Schedule')
        self.assertEqual(ws['A2'].value, u'C1 & C2')
        self.assertEqual(ws['B2'].value, u'01')          # marks stay text
        self.assertEqual(ws['C2'].value, 12.5)
        self.assertTrue(ws['A1'].font.b)
        self.assertEqual(ws.row_dimensions[2].height, 45)


if __name__ == '__main__':
    unittest.main()
